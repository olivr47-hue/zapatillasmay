# -*- coding: utf-8 -*-
"""
TikTok Shop sin API: archivos de Excel para subir a mano en Seller Center.

La conexión automática (API) de TikTok Shop exige un RFC de persona moral y el negocio es persona física, así que la
sincronización se hace con los archivos que TikTok permite subir: "Reabastecer en lote" (existencias) e "Importar en
lote" (productos nuevos). Esto antes lo hacía un script en la computadora del negocio (generar_tiktok.py) que guardaba
la llave maestra de la base de datos en un archivo; ahora corre en el servidor y la llave no sale de ahí.

Flujo:
  1. El panel sube el archivo que TikTok descarga ("Descargar todos los SKU") -> devuelve el mismo archivo con las
     existencias al día del ERP (POST /tiktok/excel/existencias).
  2. El panel sube la plantilla "Importar en lote" recién descargada -> devuelve la plantilla con SOLO los productos
     del ERP que todavía no están en TikTok (POST /tiktok/excel/nuevos). Lo que ya está en TikTok se reconoce por el SKU
     del vendedor (las primeras 3 partes del SKU: M-SAN-0002-Negro-23 -> M-SAN-0002).
"""
import base64
import io
import re
import unicodedata

from fastapi import APIRouter, File, HTTPException, UploadFile

from database import supabase_get_all

router = APIRouter(prefix="/tiktok/excel", tags=["TikTok Shop (Excel)"])

_MAX_BYTES = 12 * 1024 * 1024

# TikTok Shop exige imágenes de 100x100 a 1200x1200 px en el import masivo (las nuestras se suben a ~2000 px y las
# adicionales se rechazaban en silencio: por eso antes solo se veía la portada).
def _resize_cloudinary(url, max_px=1200):
    if not url or "res.cloudinary.com" not in url or "/upload/" not in url:
        return url
    return url.replace("/upload/", f"/upload/w_{max_px},h_{max_px},c_limit,q_auto,f_jpg/", 1)


CAT_MAP = {
    "tacones":     "Zapatos para mujer/Zapatos de tacón",
    "sandalias":   "Zapatos para mujer/Sandalias y chanclas",
    "botas":       "Zapatos para mujer/Botas",
    "botines":     "Zapatos para mujer/Botas",
    "flats":       "Zapatos para mujer/Zapatos planos",
    "plataformas": "Zapatos para mujer/Zapatos de tacón",
    "tenis":       "Zapatos para mujer/Zapatillas informales",
    "nina":        "Zapatos para mujer/Zapatillas informales",
}
CAT_DEFAULT = "Zapatos para mujer/Zapatillas informales"
SIZE_CHART_ID = "7542589447142508295"


def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise HTTPException(500, "Falta la librería openpyxl en el servidor")


def _leer_subida(archivo: UploadFile) -> bytes:
    contenido = archivo.file.read(_MAX_BYTES + 1)
    if not contenido:
        raise HTTPException(400, "El archivo está vacío")
    if len(contenido) > _MAX_BYTES:
        raise HTTPException(413, "El archivo es demasiado grande (máximo 12 MB)")
    return contenido


def _cargar_erp():
    """Productos activos + variantes + existencias del ERP, listos para cruzar con TikTok."""
    productos = supabase_get_all(
        "productos?activo=eq.true&select=id,nombre,descripcion,sku_interno,precio_menudeo,es_oferta,imagen_principal,categoria")
    variantes = supabase_get_all("variantes?activa=eq.true&select=id,producto_id,color,talla,foto_url,imagenes")
    inv = {}
    for i in supabase_get_all("inventario?select=variante_id,cantidad"):
        inv[i["variante_id"]] = inv.get(i["variante_id"], 0) + int(i.get("cantidad") or 0)
    vars_por_prod = {}
    for v in variantes:
        vars_por_prod.setdefault(v["producto_id"], []).append(v)
    inv_por_sku = {}
    _ALIAS_ID.clear()
    for p in productos:
        spu = p.get("sku_interno") or str(p["id"])
        inv_por_sku[spu] = {}
        for v in vars_por_prod.get(p["id"], []):
            color = (v.get("color") or "Unico").strip().upper()
            talla = str(v.get("talla") or "Unica").strip()
            inv_por_sku[spu][(color, talla)] = inv.get(v["id"], 0)
        if spu != str(p["id"]):
            inv_por_sku[str(p["id"])] = inv_por_sku[spu]   # mismo modelo, por si el SKU de TikTok empieza con el id
            _ALIAS_ID[str(p["id"])] = spu
    return productos, vars_por_prod, inv, inv_por_sku


_UUID_RX = re.compile(r"^[0-9a-fA-F]{8}(?:-[0-9a-fA-F]{4}){3}-[0-9a-fA-F]{12}")
_ALIAS_ID = {}   # id del producto -> sku_interno (se llena en _cargar_erp)


def _clave_color(t):
    """Color sin acentos, mayúsculas y sin espacios/guiones: «Verde metálico» = «VERDEMETALICO»."""
    s = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]", "", s)


def _clave_talla(t):
    """«25_5» (así se publicaron algunos SKU viejos en TikTok) = «25.5»."""
    v = str(t or "").strip().replace("_", ".").replace(",", ".")
    return v[:-2] if v.endswith(".0") else v


def _spu_de(sku_vend, inv_por_sku):
    """(clave del modelo en inv_por_sku, partes que quedan: color/talla). None si el modelo no está en el ERP."""
    if not sku_vend:
        return None, []
    m = _UUID_RX.match(sku_vend)
    if m:   # modelo publicado cuando aún no tenía SKU interno: el SKU empieza con su id
        spu = m.group(0).lower()
        resto = [x for x in sku_vend[m.end():].split("-") if x]
        return (spu, resto) if spu in inv_por_sku else (None, [])
    partes = sku_vend.split("-")
    if len(partes) >= 3 and "-".join(partes[:3]) in inv_por_sku:
        return "-".join(partes[:3]), partes[3:]
    if sku_vend in inv_por_sku:
        return sku_vend, []
    return None, []


def _qty_tolerante(variantes_prod, resto):
    """Cruce por color/talla ignorando acentos, espacios, abreviaturas (ROJ = ROJO, NEG = NEGRO) y 25_5 = 25.5.
    Si la abreviatura coincide con varios colores del modelo se toma la existencia MENOR (más prudente: no vender de más)."""
    idx = [(_clave_color(c), _clave_talla(t), q) for (c, t), q in variantes_prod.items()]
    tallas = {t for _, t, _ in idx}
    talla, color_tk = None, ""
    if len(resto) >= 2 and re.match(r"^\d+(?:[._,]\d+)?$", resto[-1]) and _clave_talla(resto[-1]) not in tallas:
        return None   # esa talla no existe en el ERP: no se le asigna la existencia de otra talla
    if resto and _clave_talla(resto[-1]) in tallas and len(resto) >= 2:
        talla = _clave_talla(resto[-1])
        color_tk = _clave_color("".join(resto[:-1]))
    else:
        color_tk = _clave_color("".join(resto))
    if not color_tk:
        return None
    cands = [x for x in idx if talla is None or x[1] == talla]
    exactos = [x for x in cands if x[0] == color_tk]
    if not exactos:
        exactos = [x for x in cands if x[0].startswith(color_tk) or color_tk.startswith(x[0])]
    if not exactos and len(color_tk) >= 3:
        # abreviatura por letras del color (CFO = CAFE OSCURO): solo si identifica un único color del modelo
        def _subseq(a, b):
            it = iter(b)
            return all(ch in it for ch in a)
        posibles = [x for x in cands if x[0][:1] == color_tk[:1] and _subseq(color_tk, x[0])]
        if len({x[0] for x in posibles}) == 1:
            exactos = posibles
    if not exactos:
        return None
    return min(x[2] for x in exactos) if len({x[0] for x in exactos}) > 1 else exactos[0][2]


def _buscar_qty(sku_vend, inv_por_sku):
    """Existencias del ERP para un SKU de vendedor de TikTok. None = no se encontró en el ERP."""
    spu, resto = _spu_de(sku_vend, inv_por_sku)
    if spu is None:
        return None
    variantes_prod = inv_por_sku[spu]
    if not resto:
        return sum(variantes_prod.values())
    # 1) cruce exacto (como antes)
    for n_color in range(len(resto), 0, -1):
        color = "-".join(resto[:n_color]).upper()
        talla = "-".join(resto[n_color:]) if resto[n_color:] else ""
        if (color, talla) in variantes_prod:
            return variantes_prod[(color, talla)]
        if talla == "":   # SKU solo con color: la existencia de cualquier talla de ese color
            for (c, t), qty in variantes_prod.items():
                if c == color:
                    return qty
        # (antes, si la talla no existía en el ERP, se devolvía la existencia de OTRA talla del mismo color: vender de más)
    # 2) cruce tolerante
    return _qty_tolerante(variantes_prod, resto)


def _filas_tiktok(ws):
    """[(num_fila, sku_vendedor)] del archivo de TikTok (los datos empiezan en la fila 4)."""
    filas = []
    for n, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        if row[0] is None:
            continue
        filas.append((n, str(row[3] or "").strip()))
    return filas


def _skus_en_tiktok(filas, inv_por_sku):
    """sku_interno del ERP que ya existen en TikTok (por las primeras 3 partes del SKU del vendedor)."""
    ya = set()
    for _, sv in filas:
        spu, _resto = _spu_de(sv, inv_por_sku)
        if spu is not None:
            ya.add(_ALIAS_ID.get(spu, spu))   # si el SKU empieza con el id, se cuenta como el sku_interno del modelo
    return ya


def _como_respuesta(wb, nombre, resumen):
    buf = io.BytesIO()
    wb.save(buf)
    return {"ok": True, "resumen": resumen, "nombre": nombre, "archivo_base64": base64.b64encode(buf.getvalue()).decode()}


def generar_existencias(contenido: bytes) -> dict:
    openpyxl = _openpyxl()
    try:
        wb_vista = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
        wb = openpyxl.load_workbook(io.BytesIO(contenido))
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo: sube el Excel que descarga TikTok en «Reabastecer en lote → Descargar todos los SKU»")
    ws_vista, ws = wb_vista.active, wb.active
    filas = _filas_tiktok(ws_vista)
    if not filas:
        raise HTTPException(400, "El archivo no trae productos. Asegúrate de descargar «todos los SKU» (el archivo que termina en _all_file), no el de bajo stock")
    _, _, _, inv_por_sku = _cargar_erp()
    con_stock = en_cero = sin_sku = 0
    sin_match = []
    for num, sku_vend in filas:
        qty = _buscar_qty(sku_vend, inv_por_sku)
        ws.cell(row=num, column=9, value=qty or 0)
        if not sku_vend:
            sin_sku += 1
        elif qty is None:
            sin_match.append(sku_vend)
        elif qty > 0:
            con_stock += 1
        else:
            en_cero += 1
    resumen = {"skus_en_tiktok": len(filas), "con_stock": con_stock, "agotados": en_cero,
               "sin_match_en_erp": len(sin_match), "ejemplos_sin_match": sin_match[:15], "sin_sku_vendedor": sin_sku}
    return _como_respuesta(wb, "TikTok_Actualizar_Stock.xlsx", resumen)


def generar_nuevos(contenido: bytes, export_contenido: bytes) -> dict:
    openpyxl = _openpyxl()
    productos, vars_por_prod, inv, inv_por_sku = _cargar_erp()
    # El archivo de existencias de TikTok es OBLIGATORIO: sin él no se sabe qué ya está publicado y todo saldría como
    # "nuevo" (se duplicarían los productos en TikTok).
    try:
        ws_e = openpyxl.load_workbook(io.BytesIO(export_contenido), data_only=True).active
        filas_e = _filas_tiktok(ws_e)
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo de existencias de TikTok")
    if not filas_e:
        raise HTTPException(400, "El archivo de existencias de TikTok no trae productos: descarga «todos los SKU» (el que termina en _all_file)")
    ya = _skus_en_tiktok(filas_e, inv_por_sku)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido))
    except Exception:
        raise HTTPException(400, "No se pudo leer la plantilla: descarga una nueva en «Agregar producto → Importar en lote → Descargar plantilla»")
    if "Template" not in wb.sheetnames:
        raise HTTPException(400, "Ese archivo no es la plantilla de importación de TikTok (no tiene la hoja «Template»)")
    ws = wb["Template"]
    nuevos = [p for p in productos
              if (p.get("categoria") or "").lower().strip() != "accesorios"
              and (p.get("sku_interno") or str(p["id"])) not in ya]
    if not nuevos:
        return {"ok": True, "sin_nuevos": True, "resumen": {"productos_nuevos": 0, "filas": 0, "ya_en_tiktok": len(ya)},
                "mensaje": "No hay productos nuevos: todo lo del ERP ya está en TikTok."}
    DATA_START = 6
    if ws.max_row >= DATA_START:
        ws.delete_rows(DATA_START, ws.max_row - DATA_START + 1)

    def imgs_extra(pid, principal):
        vistos = {principal} if principal else set()
        out = []
        for v in vars_por_prod.get(pid, []):
            for url in ([v.get("foto_url")] + list(v.get("imagenes") or [])):
                if url and url not in vistos:
                    vistos.add(url)
                    out.append(_resize_cloudinary(url))
                    if len(out) == 8:
                        return out
        return out

    fila, escritas = DATA_START, 0
    for p in nuevos:
        pid = p["id"]
        cat_tk = CAT_MAP.get((p.get("categoria") or "").lower().strip(), CAT_DEFAULT)
        nombre = (p.get("nombre") or p.get("sku_interno") or str(pid))[:500]
        desc = (p.get("descripcion") or nombre)[:500]
        spu = (p.get("sku_interno") or str(pid))[:200]
        # precio de la tienda (panel + $80, salvo ofertas) + $65 extra porque TikTok cobra IVA sobre el precio publicado
        precio = float(p.get("precio_menudeo") or 0) + (0 if p.get("es_oferta") else 80) + 65
        img = _resize_cloudinary(p.get("imagen_principal") or "")
        extra = (imgs_extra(pid, img) + [""] * 8)[:8]
        pvars = vars_por_prod.get(pid, [])
        if not pvars:
            filas_datos = [("Unico", "Unica", "", 0, spu)]
        else:
            filas_datos = []
            for v in pvars:
                color = (v.get("color") or "Unico").strip()[:50]
                talla = str(v.get("talla") or "Unica").strip()[:50]
                foto = _resize_cloudinary((v.get("foto_url") or "").strip()) or img
                filas_datos.append((color, talla, foto, inv.get(v["id"], 0), f"{spu}-{color}-{talla}"[:200]))
        for color, talla, foto, qty, sku_var in filas_datos:
            datos = [
                cat_tk, "", nombre, desc, img,
                extra[0], extra[1], extra[2], extra[3], extra[4], extra[5], extra[6], extra[7],
                "", "",
                "Color", color, foto,
                "Talla", talla,
                1000, 30, 20, 10, "",
                precio, qty, sku_var, 1, "", "",
                SIZE_CHART_ID,
                "", "", "", "", "", "", "",
            ]
            for col, valor in enumerate(datos, 1):
                ws.cell(row=fila, column=col, value=valor)
            fila += 1
            escritas += 1
    resumen = {"productos_nuevos": len(nuevos), "filas": escritas, "ya_en_tiktok": len(ya),
               "ejemplos": [(p.get("nombre") or p.get("sku_interno") or "")[:60] for p in nuevos[:8]]}
    return _como_respuesta(wb, "TikTok_Productos_Nuevos.xlsx", resumen)


@router.post("/existencias")
def excel_existencias(archivo: UploadFile = File(...)):
    """Recibe el Excel de «Reabastecer en lote → Descargar todos los SKU» y lo devuelve con las existencias del ERP."""
    return generar_existencias(_leer_subida(archivo))


@router.post("/nuevos")
def excel_nuevos(plantilla: UploadFile = File(...), existencias: UploadFile = File(...)):
    """Recibe la plantilla de «Importar en lote» y el Excel de existencias de TikTok (para saber qué ya está publicado) y
    devuelve la plantilla llena con los productos que todavía no están en TikTok."""
    return generar_nuevos(_leer_subida(plantilla), _leer_subida(existencias))
